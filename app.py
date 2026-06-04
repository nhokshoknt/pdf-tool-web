import os
import re
import pdfplumber
from flask import Flask, request, render_template, send_file
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
OUTPUT_FOLDER = "outputs"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

excel_path = os.path.join(OUTPUT_FOLDER, "ket_qua.xlsx")


# =========================
# CLEAN TEXT
# =========================
def clean_text(value):
    if not value:
        return ""

    text = value.group(1)
    text = text.replace("\n", " ")
    text = text.replace("Tỉnh Khánh Hòa", "")

    text = re.sub(r",\s*Nha Trang.*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^,\s*", "", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def clean_filename(name):
    name = re.sub(r'[\\/*?:"<>|]', "", name)
    name = re.sub(r"\s+", " ", name)
    return name.strip()


# =========================
# PROCESS PDF
# =========================
def process_pdf(file_path):
    text = ""

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"

    name_match = re.search(r"Tên KH:\s*(.*?)\s*CMND", text, re.IGNORECASE | re.DOTALL)

    if not name_match:
        return None

    tamtru_match = re.search(r"Đ/c Tạm trú:\s*(.+?)(?:Đ/c Hộ khẩu:)", text, re.IGNORECASE | re.DOTALL)
    hokhau_match = re.search(r"Đ/c Hộ khẩu:\s*(.+?)(?:Đ/c Làm việc:)", text, re.IGNORECASE | re.DOTALL)
    lamviec_match = re.search(r"Đ/c Làm việc:\s*(.+?)(?:THÔNG TIN KHOẢN VAY|Số HĐ:)", text, re.IGNORECASE | re.DOTALL)

    name = name_match.group(1).replace("\n", " ").strip()

    hokhau = clean_text(hokhau_match)
    tamtru = clean_text(tamtru_match)
    lamviec = clean_text(lamviec_match)

    return name, hokhau, tamtru, lamviec


# =========================
# INIT EXCEL
# =========================
def init_excel():
    if os.path.exists(excel_path):
        wb = load_workbook(excel_path)
        ws = wb.active
        stt = ws.max_row
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["STT", "Tên KH", "Đ/c Hộ khẩu", "Đ/c Tạm trú", "Đ/c Làm việc"])
        stt = 1

    return wb, ws, stt


# =========================
# ROUTE
# =========================
@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        files = request.files.getlist("files")

        wb, ws, stt = init_excel()

        for file in files:
            filename = secure_filename(file.filename)
            path = os.path.join(UPLOAD_FOLDER, filename)
            file.save(path)

            result = process_pdf(path)

            if not result:
                continue

            name, hokhau, tamtru, lamviec = result

            # check duplicate
            is_duplicate = False
            for row in ws.iter_rows(min_row=2, values_only=True):
                if (
                    name == row[1]
                    and hokhau == row[2]
                    and tamtru == row[3]
                    and lamviec == row[4]
                ):
                    is_duplicate = True
                    break

            if not is_duplicate:
                ws.append([stt, name, hokhau, tamtru, lamviec])
                stt += 1

            # rename file
            new_name = clean_filename(name) + ".pdf"
            new_path = os.path.join(UPLOAD_FOLDER, new_name)

            os.rename(path, new_path)

        wb.save(excel_path)

        return send_file(excel_path, as_attachment=True)

    return render_template("index.html")

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)